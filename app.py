import os
import io
import re
import json
import difflib
import base64
import html as _html
from typing import List, Tuple, Dict

import streamlit as st
try:
    from streamlit_lottie import st_lottie
except Exception:  # pragma: no cover
    st_lottie = None
try:
    from streamlit_toggle import st_toggle_switch
except Exception:  # pragma: no cover
    st_toggle_switch = None
from textstat import textstat
from fpdf import FPDF
from rake_nltk import Rake
import nltk
import requests
from bs4 import BeautifulSoup
from PyPDF2 import PdfReader
from docx import Document

try:
    import language_tool_python
except Exception:  # pragma: no cover
    language_tool_python = None


# -----------------------------
# App Config
# -----------------------------
st.set_page_config(
    page_title="Medindia Content Optimizer",
    page_icon="🩺",
    layout="wide",
    initial_sidebar_state="expanded",
)


# -----------------------------
# Utilities
# -----------------------------
@st.cache_data(show_spinner=False)
def _ensure_nltk():
    try:
        nltk.data.find('tokenizers/punkt')
    except LookupError:
        nltk.download('punkt', quiet=True)
    try:
        nltk.data.find('corpora/stopwords')
    except LookupError:
        nltk.download('stopwords', quiet=True)


@st.cache_data(show_spinner=False)
def load_lottie_url(url: str):
    try:
        r = requests.get(url, timeout=10)
        if r.status_code == 200:
            return r.json()
    except Exception:
        return None
    return None


@st.cache_data(show_spinner=False)
def fetch_url_text(url: str) -> str:
    headers = {"User-Agent": "Mozilla/5.0 (compatible; MedindiaBot/1.0)"}
    resp = requests.get(url, headers=headers, timeout=15)
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, 'html.parser')
    # Remove script/style
    for t in soup(["script", "style", "nav", "footer", "header", "noscript"]):
        t.decompose()
    # Medindia main content heuristic
    article = soup.find('article') or soup.find('main') or soup
    text = article.get_text(separator=" ", strip=True)
    # Basic cleanup
    return re.sub(r"\s+", " ", text)


@st.cache_data(show_spinner=False)
def parse_uploaded_file(file) -> str:
    name = (file.name or "").lower()
    if name.endswith(".txt"):
        return file.read().decode('utf-8', errors='ignore')
    if name.endswith(".pdf"):
        reader = PdfReader(file)
        pages = [p.extract_text() or "" for p in reader.pages]
        return "\n\n".join(pages)
    if name.endswith(".docx"):
        doc = Document(file)
        return "\n".join(p.text for p in doc.paragraphs)
    return ""


@st.cache_data(show_spinner=False)
def run_language_tool(text: str) -> Tuple[str, List[dict]]:
    if not text.strip():
        return text, []
    if language_tool_python is None:
        # Fallback: return original if library unavailable
        return text, []
    tool = language_tool_python.LanguageToolPublicAPI('en-US')
    corrected = tool.correct(text)
    matches = tool.check(text)
    # Convert matches to serializable dicts (subset)
    simple_matches = []
    for m in matches:
        simple_matches.append({
            "offset": m.offset,
            "errorLength": m.errorLength,
            "message": m.message,
            "replacements": [r.value for r in m.replacements][:3],
            "ruleId": m.ruleId,
        })
    return corrected, simple_matches


def _escape_html(text: str) -> str:
    return _html.escape(text, quote=True)


def highlight_diff(original: str, corrected: str) -> str:
    def tokenize(s: str) -> List[str]:
        return re.findall(r"\w+|\s+|[\.,!?;:()\-]", s)

    o_tokens = tokenize(original)
    c_tokens = tokenize(corrected)
    diff = difflib.SequenceMatcher(a=o_tokens, b=c_tokens)
    html_parts: List[str] = []
    for tag, i1, i2, j1, j2 in diff.get_opcodes():
        if tag == 'equal':
            html_parts.append("".join(c_tokens[j1:j2]))
        elif tag == 'replace':
            del_text = "".join(o_tokens[i1:i2])
            add_text = "".join(c_tokens[j1:j2])
            if del_text.strip():
                html_parts.append(f'<span class="change del" title="Removed">{_escape_html(del_text)}</span>')
            if add_text.strip():
                html_parts.append(f'<span class="change add" title="Added">{_escape_html(add_text)}</span>')
        elif tag == 'delete':
            del_text = "".join(o_tokens[i1:i2])
            if del_text.strip():
                html_parts.append(f'<span class="change del" title="Removed">{_escape_html(del_text)}</span>')
        elif tag == 'insert':
            add_text = "".join(c_tokens[j1:j2])
            if add_text.strip():
                html_parts.append(f'<span class="change add" title="Added">{_escape_html(add_text)}</span>')
    return "".join(html_parts)


# Simple medical term dictionary (mock). Extend as needed.
MEDICAL_TERMS: Dict[str, Dict[str, List[str] | str]] = {
    "endometriosis": {
        "definition": "A condition where tissue similar to the uterine lining grows outside the uterus.",
        "symptoms": ["Pelvic pain", "Heavy periods", "Infertility"],
        "causes": ["Retrograde menstruation", "Immune disorders"],
        "treatments": ["Pain relievers", "Hormone therapy", "Surgery"],
    },
    "hypertension": {
        "definition": "Abnormally high blood pressure that increases cardiovascular risk.",
        "symptoms": ["Often none", "Headaches", "Shortness of breath"],
        "causes": ["Genetics", "Diet", "Lifestyle"],
        "treatments": ["Lifestyle changes", "ACE inhibitors", "Diuretics"],
    },
}


def annotate_medical_terms(html_text: str) -> str:
    def replacement(match):
        term = match.group(0)
        key = term.lower()
        if key in MEDICAL_TERMS:
            info = MEDICAL_TERMS[key]
            tooltip = (
                f"<strong>{term.title()}</strong><br>"
                f"{_escape_html(info['definition'])}<br>"
                f"<em>Symptoms</em>: " + ", ".join(info['symptoms']) + "<br>"
                f"<em>Treatments</em>: " + ", ".join(info['treatments'])
            )
            return f'<span class="med-term" data-tip="{_escape_html(tooltip)}">{_escape_html(term)}</span>'
        return term

    pattern = re.compile(r"\b(" + "|".join(map(re.escape, MEDICAL_TERMS.keys())) + r")\b", re.IGNORECASE)
    return pattern.sub(replacement, html_text)


def seo_keywords(text: str) -> List[str]:
    _ensure_nltk()
    r = Rake()  # uses stopwords and punctuation
    r.extract_keywords_from_text(text)
    phrases = r.get_ranked_phrases()[:15]
    # Split long phrases into keywords if needed
    keywords: List[str] = []
    for p in phrases:
        if len(p.split()) <= 3:
            keywords.append(p)
        else:
            keywords.extend([w for w in p.split() if len(w) > 3])
    # Deduplicate preserve order
    seen = set()
    out = []
    for k in keywords:
        k2 = k.lower()
        if k2 not in seen:
            seen.add(k2)
            out.append(k)
    return out[:20]


def compute_seo_score(text: str, keywords: List[str]) -> int:
    # Very simple heuristic combining length, keyword richness, readability
    word_count = len(re.findall(r"\w+", text))
    readability = max(0, min(100, 100 - textstat.difficult_words(text)))
    richness = min(100, int(len(keywords) * 5))
    length_score = min(100, int((word_count / 1200) * 100))  # target ~1200 words
    score = int(0.4 * readability + 0.3 * richness + 0.3 * length_score)
    return max(0, min(100, score))


def make_pdf(content: str, title: str = "Processed Article") -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=14)
    pdf.cell(0, 10, txt=title, ln=True, align='C')
    pdf.set_font("Helvetica", size=11)
    for line in content.splitlines():
        pdf.multi_cell(0, 6, txt=line)
    return pdf.output(dest='S').encode('latin1', errors='ignore')


def load_css(path: str):
    if os.path.exists(path):
        with open(path, 'r', encoding='utf-8') as f:
            st.markdown(f"<style>{f.read()}</style>", unsafe_allow_html=True)


# -----------------------------
# Session State Defaults
# -----------------------------
if "dark_mode" not in st.session_state:
    st.session_state.dark_mode = True
if "processed" not in st.session_state:
    st.session_state.processed = False
if "corrected_text" not in st.session_state:
    st.session_state.corrected_text = ""
if "original_text" not in st.session_state:
    st.session_state.original_text = ""
if "email_to" not in st.session_state:
    st.session_state.email_to = ""
if "email_subject" not in st.session_state:
    st.session_state.email_subject = "Processed Health Article"
if "email_cc" not in st.session_state:
    st.session_state.email_cc = ""
if "email_bcc" not in st.session_state:
    st.session_state.email_bcc = ""
if "email_body" not in st.session_state:
    st.session_state.email_body = ""


# -----------------------------
# Assets & Styles
# -----------------------------
load_css("styles.css")

# Gradient background and intro animation
st.markdown(
    """
    <div class="gradient-bg"></div>
    <div class="fade-in"></div>
    """,
    unsafe_allow_html=True,
)


# -----------------------------
# Header with controls
# -----------------------------
col_title, col_actions = st.columns([0.7, 0.3], gap="large")
with col_title:
    title_clicked = st.button("MedIndia4u Pvt Ltd", key="brand", help="Click to refresh")
    if title_clicked:
        st.rerun()
    st.caption("Premium Health Content Proofreader ✨")

with col_actions:
    # Dark mode toggle via streamlit-toggle-switch, fallback to st.toggle
    if st_toggle_switch is not None:
        st.session_state.dark_mode = st_toggle_switch(
            label="Dark Mode",
            default_value=st.session_state.dark_mode,
            label_after=True,
            inactive_color="#9CA3AF",
            active_color="#10B981",
            track_color="#374151",
            key="dark_toggle",
        )
    else:
        st.session_state.dark_mode = st.toggle("Dark Mode", value=st.session_state.dark_mode)

    share_clicked = st.button("Share ✉️", use_container_width=True)
    st.markdown("<div class='avatar'>🧑‍⚕️</div>", unsafe_allow_html=True)

if st.session_state.dark_mode:
    st.markdown("<style>body, .stApp { color: #E5E7EB !important; }</style>", unsafe_allow_html=True)
if share_clicked:
    st.session_state.show_email = True


# -----------------------------
# Hero Section
# -----------------------------
hero_col1, hero_col2 = st.columns([0.6, 0.4])
with hero_col1:
    st.markdown(
        "<div class='hero-card glass'>🩺 <span class='hero-title'>Optimize Your Health Content Effortlessly</span><br><span class='hero-sub'>Notion meets Grammarly — tuned for medical accuracy.</span></div>",
        unsafe_allow_html=True,
    )
with hero_col2:
    lottie = load_lottie_url("https://assets3.lottiefiles.com/packages/lf20_jcikwtux.json")
    if st_lottie and lottie:
        st_lottie(lottie, height=160, key="heartbeat")
    else:
        st.markdown("<div class='glass' style='text-align:center;'>💜💙</div>", unsafe_allow_html=True)


# -----------------------------
# Sidebar: Input Section
# -----------------------------
with st.sidebar:
    st.markdown("## Input Your Content")
    with st.form("input_form"):
        pasted = st.text_area("Paste Text", placeholder="Paste your article here...", height=180)
        url = st.text_input("Paste URL", placeholder="e.g., https://www.medindia.net/...")
        uploaded = st.file_uploader("Upload Document", type=["txt", "docx", "pdf"], accept_multiple_files=False)
        submitted = st.form_submit_button("Process", use_container_width=True)

    if submitted:
        with st.spinner("Processing your content..."):
            text = ""
            if pasted and pasted.strip():
                text = pasted
            elif url and url.strip():
                try:
                    text = fetch_url_text(url.strip())
                except Exception as e:
                    st.error(f"Failed to fetch URL: {e}")
            elif uploaded is not None:
                try:
                    text = parse_uploaded_file(uploaded)
                except Exception as e:
                    st.error(f"Failed to parse file: {e}")
            else:
                st.warning("Please provide text, a URL, or upload a document.")

            if text and text.strip():
                st.session_state.original_text = text
                corrected, matches = run_language_tool(text)
                st.session_state.corrected_text = corrected
                st.session_state.processed = True
            else:
                st.session_state.processed = False


# -----------------------------
# Main Output
# -----------------------------
st.markdown("### Proofed & Optimized Article")
if st.session_state.processed and st.session_state.corrected_text:
    original = st.session_state.original_text
    corrected = st.session_state.corrected_text

    diff_html = highlight_diff(original, corrected)
    diff_html = annotate_medical_terms(diff_html)

    with st.container(border=True):
        st.markdown("<div class='article-card glass'>" + diff_html + "</div>", unsafe_allow_html=True)

    # Stats
    wc = len(re.findall(r"\w+", corrected))
    readability = textstat.flesch_reading_ease(corrected)
    keywords = seo_keywords(corrected)
    seo_score = compute_seo_score(corrected, keywords)

    with st.expander("SEO Suggestions", expanded=False):
        # Simple meta title suggestion
        top_kw = ", ".join(keywords[:4]) if keywords else "Health Article"
        title_suggestion = f"{top_kw} | Medindia"
        meta_desc = corrected.strip().split(". ")
        meta_desc = (meta_desc[0] if meta_desc else "Optimized health content by Medindia.")
        st.markdown(f"**Title**: {title_suggestion}")
        st.markdown(f"**Meta Description**: {meta_desc[:160]}")

    c1, c2, c3, c4 = st.columns([1, 1, 1, 2])
    with c1:
        st.metric("Words", wc)
    with c2:
        st.metric("Readability", f"{readability:.1f}")
    with c3:
        st.metric("SEO Score", f"{seo_score}/100")
    with c4:
        st.markdown("**Suggested Keywords**")
        st.markdown(
            " ".join([f"<span class='tag'>{_escape_html(k)}</span>" for k in keywords[:12]]),
            unsafe_allow_html=True,
        )
    st.progress(seo_score / 100)

    # Actions
    dl_col, pdf_col, share_col = st.columns(3)
    with dl_col:
        st.download_button("Copy / Download TXT", corrected, file_name="processed_article.txt")
    with pdf_col:
        pdf_bytes = make_pdf(corrected)
        st.download_button("Export PDF", data=pdf_bytes, file_name="processed_article.pdf", mime="application/pdf")
    with share_col:
        if st.button("Send Processed Content ✉️", use_container_width=True):
            st.session_state.show_email = True

    if getattr(st.session_state, 'show_email', False):
        with st.expander("Email Composer", expanded=True):
            preset = st.button("Use Medindia Team Preset")
            if preset:
                st.session_state.email_to = "editor@medindia.net"
                st.session_state.email_cc = "team@medindia.net"
                st.session_state.email_subject = "Medindia | Proofed Article"
                st.session_state.email_body = corrected

            to = st.text_input("To", key="email_to", placeholder="editor@medindia.net")
            subject = st.text_input("Subject", key="email_subject")
            body = st.text_area("Body", key="email_body", value=st.session_state.get("email_body", corrected) or corrected, height=180)
            cc = st.text_input("CC", key="email_cc", placeholder="team@medindia.net")
            bcc = st.text_input("BCC", key="email_bcc", placeholder="")
            # mailto link (with cc/bcc if provided)
            params = []
            params.append(f"subject={requests.utils.requote_uri(subject)}")
            params.append(f"body={requests.utils.requote_uri((body or corrected)[:2000])}")
            if cc:
                params.append(f"cc={requests.utils.requote_uri(cc)}")
            if bcc:
                params.append(f"bcc={requests.utils.requote_uri(bcc)}")
            query = "&".join(params)
            mailto = "mailto:" + requests.utils.requote_uri(to or "") + (f"?{query}" if query else "")
            st.link_button("Open in Email Client ✉️", url=mailto)
            st.caption("Uses a simple mailto link. No SMTP setup required.")
else:
    st.info("Submit content via the sidebar to see the optimized article here.")


# -----------------------------
# Footer Contact Card
# -----------------------------
st.markdown("---")
with st.container():
    st.markdown(
        """
        <div class="footer-card glass">
            <div class="footer-title">📍 Contact Address</div>
            <div>India<br>Medindia4u.com Pvt. Ltd.,<br>1st Floor, New No.10, Old No.6,<br>AE-Block, 7th Street, 10th Main Road,<br>Anna Nagar,<br>Chennai-600 040<br>India<br>Tel: <a href="tel:+919791173039">+91 97911 73039</a><br>Email: <a href="mailto:info@medindia.net">info@medindia.net</a></div>
            <div class="socials">🔗 <a href="https://www.medindia.net" target="_blank" rel="noreferrer">Website</a> · 🐦 <a href="https://x.com" target="_blank" rel="noreferrer">Twitter</a> · 🔵 <a href="https://www.linkedin.com" target="_blank" rel="noreferrer">LinkedIn</a></div>
        </div>
        """,
        unsafe_allow_html=True,
    )


# -----------------------------
# Inline CSS fallbacks (if styles.css not found)
# -----------------------------
if not os.path.exists("styles.css"):
    st.markdown(
        """
        <style>
        .gradient-bg { position: fixed; inset: 0; background: linear-gradient(135deg,#111827, #1f2937, #0ea5e9, #8b5cf6) fixed; opacity: 0.18; z-index: -1; }
        .fade-in { animation: fadeIn 0.8s ease-in-out; }
        @keyframes fadeIn { from{opacity:0; transform: translateY(6px);} to{opacity:1; transform:none;} }
        .glass { background: rgba(255,255,255,0.06); border: 1px solid rgba(255,255,255,0.12); border-radius: 16px; backdrop-filter: blur(8px); padding: 16px; }
        .hero-card { font-size: 18px; }
        .hero-title { font-weight: 700; font-size: 26px; }
        .hero-sub { color: #9CA3AF; }
        .avatar { text-align: right; font-size: 28px; margin-top: 6px; }
        .article-card { line-height: 1.7; font-size: 1.02rem; }
        .change.add { background: rgba(16,185,129,0.18); border-bottom: 2px solid #10B981; }
        .change.del { background: rgba(239,68,68,0.15); text-decoration: line-through; }
        .med-term { position: relative; border-bottom: 1px dashed #60A5FA; cursor: help; }
        .med-term:hover::after { content: attr(data-tip); position: absolute; left: 0; top: 1.6em; white-space: normal; background: #111827; color: #E5E7EB; border: 1px solid #374151; padding: 10px 12px; border-radius: 10px; width: 280px; z-index: 10; }
        .tag { display: inline-block; padding: 4px 10px; margin: 3px; background: #111827; color: #93C5FD; border-radius: 999px; font-size: 0.85rem; }
        .footer-card { margin-top: 24px; }
        .footer-title { font-weight: 700; margin-bottom: 8px; }
        </style>
        """,
        unsafe_allow_html=True,
    )


